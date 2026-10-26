"""Actual extraction/raster tests plus explicit isolated failure-injection tests.

PaddleOCR recognition is NOT mocked into an acceptance pass: the local sample
must report blocked/partial when the real models/provider are not configured.
"""
import base64
import contextlib
import hashlib
import importlib.util
import io
import json
import os
from pathlib import Path
import socket
import subprocess
import sys
import tempfile
import threading
import time
import unittest
from unittest import mock
import zipfile

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import adapter
from PIL import Image, ImageDraw
from pypdf import PdfReader, PdfWriter
from reportlab.pdfgen import canvas


def request(data=b"Hello CloudDrive\n", filename="note.txt", **recipe):
    return {"job_id": "test-job", "attempt_id": "1", "input": {"filename": filename, "content_base64": base64.b64encode(data).decode(), "sha256": hashlib.sha256(data).hexdigest()}, "recipe": recipe, "output_scope": "tmp/test-job/1/"}


def pdf(pages=1, scanned=()):
    out = io.BytesIO()
    doc = canvas.Canvas(out)
    for page in range(1, pages + 1):
        if page in scanned:
            # Real image-only PDF, no hidden PDF text layer.
            from reportlab.lib.utils import ImageReader
            image = Image.new("RGB", (800, 200), "white")
            ImageDraw.Draw(image).text((20, 60), f"Scanned document page {page}", fill="black")
            doc.drawImage(ImageReader(image), 50, 550, width=400, height=100)
        else:
            doc.drawString(50, 750, f"Extractable CloudDrive page {page}. Stable source version.")
        doc.showPage()
    doc.save()
    return out.getvalue()


class SecurityTests(unittest.TestCase):
    def test_path_escape_output(self):
        r = request(); r["output_scope"] = "tmp/test-job/../1/"
        self.assertEqual(adapter.execute(r)["errors"][0]["code"], "invalid_output_scope")

    def test_bad_job_identifier(self):
        r = request(); r["job_id"] = "../outside"
        self.assertEqual(adapter.execute(r)["status"], "failed")

    def test_checksum(self):
        r = request(); r["input"]["sha256"] = "0" * 64
        self.assertEqual(adapter.execute(r)["errors"][0]["code"], "checksum_mismatch")

    def test_unsupported(self):
        self.assertEqual(adapter.execute(request(b"raw", "disk.exe"))["status"], "not_converted")

    def test_filename_escape(self):
        self.assertEqual(adapter.execute(request(b"raw", "../secret.txt"))["errors"][0]["code"], "invalid_filename")

    def test_url_ssrf(self):
        for url in ("file:///etc/passwd", "http://169.254.169.254/latest/meta-data/", "http://api:8080@evil.test/internal/source/a", "http://api:8080/internal/source/../admin", "http://api:8080/internal/source/%2e%2e/admin", "http://api:8080/internal/source/%252e%252e/admin", "http://api:8080.evil.test/internal/source/a"):
            with self.subTest(url=url), self.assertRaises(adapter.AdapterError):
                adapter.checked_url(url, ["http://api:8080/internal/source/"])
        self.assertEqual(adapter.checked_url("http://api:8080/internal/source/abc?ticket=foo", ["http://api:8080/internal/source/"]), "http://api:8080/internal/source/abc?ticket=foo")

    def test_redirect_rejected(self):
        with self.assertRaises(adapter.AdapterError):
            adapter.NoRedirect().redirect_request(None, None, 302, "", {}, "http://evil.test/")

    def test_request_cannot_increase_limit(self):
        self.assertEqual(adapter.execute(request(max_pages=adapter.MAX_PAGES + 1))["status"], "not_converted")

    def test_zip_expansion_limit(self):
        out = io.BytesIO()
        with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as archive:
            archive.writestr("a", b"A" * 1000000)
        self.assertEqual(adapter.execute(request(out.getvalue(), "bomb.docx"))["errors"][0]["code"], "resource_limit")

    def test_token_and_source_not_echoed(self):
        r = request(); del r["input"]["content_base64"]
        r["input"]["url"] = "http://evil.invalid/secrets?ticket=do-not-leak"
        body = json.dumps(adapter.execute(r))
        self.assertNotIn("do-not-leak", body)
        self.assertNotIn("evil.invalid", body)


class FormatTests(unittest.TestCase):
    def test_utf8_markdown(self):
        text = "# 工程报告\n\n版本固定和历史保留。"
        result = adapter.execute(request(text.encode(), "notes.md"))
        self.assertEqual(result["status"], "ready")
        self.assertEqual(result["markdown"], text)

    def test_invalid_utf8(self):
        self.assertEqual(adapter.execute(request(b"\xff", "bad.txt"))["errors"][0]["code"], "malformed")

    def test_text_pdf(self):
        result = adapter.execute(request(pdf(3), "sample.pdf"))
        self.assertEqual(result["status"], "ready", result)
        self.assertEqual(result["completeness"]["total_pages"], 3)
        self.assertIn("Extractable CloudDrive page 3", result["markdown"])
        self.assertEqual([p["page"] for p in result["pages"]], [1, 2, 3])

    def test_encrypted_pdf(self):
        writer = PdfWriter(); writer.add_page(PdfReader(io.BytesIO(pdf())).pages[0]); writer.encrypt("secret")
        data = io.BytesIO(); writer.write(data)
        self.assertEqual(adapter.execute(request(data.getvalue(), "secret.pdf"))["errors"][0]["code"], "encrypted")

    def test_pdf_page_limit(self):
        self.assertEqual(adapter.execute(request(pdf(3), "sample.pdf", max_pages=2))["status"], "not_converted")

    def test_scan_no_fake_success(self):
        with mock.patch.dict(os.environ, {"DRIVE_OCR_URL": "", "DRIVE_VISION_BASE_URL": "", "DRIVE_VISION_MODEL": ""}):
            result = adapter.execute(request(pdf(3, scanned=(2,)), "scan.pdf"))
        self.assertEqual(result["status"], "partial")
        self.assertEqual(result["completeness"]["missing_text_pages"], [2])
        self.assertEqual(result["completeness"]["missing_description_pages"], [2])
        self.assertIn("page 3", result["markdown"])

    def test_200_page_failure_mapping_only_not_ocr_acceptance(self):
        missing = (17, 48, 91, 130, 199)
        # The entire 200-page PDF is real; absent providers deliberately fail the
        # five image-only pages. This validates accounting, NOT OCR recognition.
        with mock.patch.dict(os.environ, {"DRIVE_OCR_URL": "", "DRIVE_VISION_BASE_URL": "", "DRIVE_VISION_MODEL": ""}):
            result = adapter.execute(request(pdf(200, scanned=missing), "long.pdf"))
        self.assertEqual(result["status"], "partial")
        self.assertEqual(result["completeness"]["missing_text_pages"], list(missing))
        self.assertEqual(len(result["pages"]), 200)
        self.assertIn("page 200", result["markdown"])

    def test_scan_page_marker_precedes_assets_and_description(self):
        with mock.patch("adapter.ocr_remote", side_effect=adapter.AdapterError("ocr_failed", "ocr")), mock.patch("adapter.describe", return_value="Visible diagram belongs to page two"):
            result = adapter.execute(request(pdf(3, scanned=(2,)), "scan.pdf"))
        markdown = result["markdown"]
        marker = markdown.index("<!-- page:2 block:p2-b01 -->")
        self.assertLess(marker, markdown.index("![Page 2]"))
        self.assertLess(marker, markdown.index("Visible diagram belongs to page two"))
        self.assertLess(markdown.index("Visible diagram belongs to page two"), markdown.index("<!-- page:3"))
        self.assertEqual(markdown.count("<!-- page:2"), 1)

    def test_selected_page_preserves_global_number(self):
        result = adapter.execute(request(pdf(3), "sample.pdf", page=2))
        self.assertEqual(result["pages"][0]["page"], 2)
        self.assertEqual(result["completeness"]["total_pages"], 3)
        self.assertNotIn("page 1", result["markdown"])

    def test_pdf_preview_only_png(self):
        result = adapter.execute(request(pdf(2), "sample.pdf"), mode="preview")
        self.assertEqual(result["status"], "ready")
        self.assertEqual(result["markdown"], "")
        self.assertEqual(len(result["assets"]), 2)
        for item in result["assets"]:
            body = base64.b64decode(item["content_base64"])
            self.assertEqual(item["media_type"], "image/png")
            self.assertTrue(body.startswith(b"\x89PNG"))
            self.assertEqual(hashlib.sha256(body).hexdigest(), item["sha256"])

    def test_image_reencoded_metadata_stripped(self):
        from PIL.PngImagePlugin import PngInfo
        meta = PngInfo(); meta.add_text("private", "do-not-leak")
        out = io.BytesIO(); Image.new("RGB", (50, 50), "red").save(out, "PNG", pnginfo=meta)
        result = adapter.execute(request(out.getvalue(), "photo.png"), mode="preview")
        body = base64.b64decode(result["assets"][0]["content_base64"])
        self.assertNotIn(b"do-not-leak", body)

    def test_svg_fails_closed(self):
        result = adapter.execute(request(b'<svg xmlns="http://www.w3.org/2000/svg"><script>alert(1)</script></svg>', "attack.svg"), mode="preview")
        self.assertEqual(result["status"], "not_converted")
        self.assertEqual(result["assets"], [])

    def test_static_svg_raster(self):
        result = adapter.execute(request(b'<svg xmlns="http://www.w3.org/2000/svg" width="100" height="100"><rect width="100" height="100" fill="red"/></svg>', "safe.svg"), mode="preview")
        self.assertEqual(result["status"], "ready", result)
        self.assertTrue(base64.b64decode(result["assets"][0]["content_base64"]).startswith(b"\x89PNG"))

    def test_svg_external_fetch_forbidden(self):
        for xml in (b'<svg xmlns="http://www.w3.org/2000/svg"><use href="http://169.254.169.254/latest/meta-data"/></svg>', b'<svg xmlns="http://www.w3.org/2000/svg"><style>@import "https://evil.invalid";</style></svg>'):
            result = adapter.execute(request(xml, "unsafe.svg"), mode="preview")
            self.assertEqual(result["status"], "not_converted", result)
            self.assertEqual(result["assets"], [])

    def test_real_paddle_unconfigured_is_not_pass(self):
        out = io.BytesIO(); Image.new("RGB", (30, 30), "white").save(out, "PNG")
        with tempfile.TemporaryDirectory() as path, mock.patch.dict(os.environ, {"PADDLE_MODELS_DIR": path}):
            result = adapter.execute(request(out.getvalue(), "scan.png"), mode="paddleocr")
        self.assertEqual(result["status"], "blocked")
        self.assertEqual(result["errors"][0]["code"], "provider_not_configured")


@unittest.skipUnless(importlib.util.find_spec("anydoc"), "real pinned upstream AnyDoc wheel not installed")
class ActualAnyDocTests(unittest.TestCase):
    def test_csv_upstream(self):
        result = adapter.execute(request(b"name,value\nAlice,17\n", "table.csv"))
        self.assertEqual(result["status"], "ready", result)
        self.assertEqual(result["engine"], {"name": "firecrawl-anydoc", "version": "0.2.4"})
        self.assertIn("| Alice | 17 |", result["markdown"])

    def test_rtf_upstream(self):
        result = adapter.execute(request(br"{\rtf1\ansi CloudDrive RTF document.}", "document.rtf"))
        self.assertEqual(result["status"], "ready", result)
        self.assertIn("CloudDrive RTF", result["markdown"])

    def test_docx_upstream(self):
        from docx import Document
        document = Document(); document.add_heading("CloudDrive specification", 0); document.add_paragraph("Private document version 7.")
        out = io.BytesIO(); document.save(out)
        result = adapter.execute(request(out.getvalue(), "document.docx"))
        self.assertEqual(result["status"], "ready", result)
        self.assertIn("Private document version 7", result["markdown"])

    def test_xlsx_upstream(self):
        from openpyxl import Workbook
        workbook = Workbook(); workbook.active.append(["Owner", "Quota"]); workbook.active.append(["Alice", 200])
        out = io.BytesIO(); workbook.save(out)
        result = adapter.execute(request(out.getvalue(), "table.xlsx"))
        self.assertEqual(result["status"], "ready", result)
        self.assertIn("Alice", result["markdown"])
        self.assertIn("200", result["markdown"])

    def test_pptx_upstream(self):
        from pptx import Presentation
        presentation = Presentation(); slide = presentation.slides.add_slide(presentation.slide_layouts[1]); slide.shapes.title.text = "CloudDrive slides"; slide.placeholders[1].text = "Stable version citations"
        out = io.BytesIO(); presentation.save(out)
        result = adapter.execute(request(out.getvalue(), "slides.pptx"))
        self.assertEqual(result["status"], "ready", result)
        self.assertIn("Stable version citations", result["markdown"])

    def test_xlsx_native_chart_is_not_silently_complete(self):
        from openpyxl import Workbook
        from openpyxl.chart import BarChart, Reference
        workbook = Workbook(); sheet = workbook.active
        sheet.append(["Quarter", "Revenue"]); sheet.append(["Q1", 100]); sheet.append(["Q2", 200])
        chart = BarChart(); chart.add_data(Reference(sheet, min_col=2, min_row=1, max_row=3), titles_from_data=True)
        sheet.add_chart(chart, "D1")
        out = io.BytesIO(); workbook.save(out)
        result = adapter.execute(request(out.getvalue(), "chart.xlsx"))
        self.assertEqual(result["status"], "partial", result)
        self.assertIn("Revenue", result["markdown"])
        self.assertEqual(result["completeness"]["missing_description_pages"], [1])
        self.assertIn("native_visual_object_not_rendered", [error["code"] for error in result["errors"]])

    def test_docx_embedded_image_preserved_but_missing_description(self):
        from docx import Document
        image = io.BytesIO(); Image.new("RGB", (100, 100), "blue").save(image, "PNG"); image.seek(0)
        document = Document(); document.add_paragraph("Embedded diagram."); document.add_picture(image)
        out = io.BytesIO(); document.save(out)
        with mock.patch.dict(os.environ, {"DRIVE_VISION_BASE_URL": "", "DRIVE_VISION_MODEL": ""}):
            result = adapter.execute(request(out.getvalue(), "images.docx"))
        self.assertEqual(result["status"], "partial", result)
        self.assertEqual(len(result["assets"]), 1, result)
        self.assertEqual(result["completeness"]["missing_description_pages"], [1])


class HTTPTests(unittest.TestCase):
    def test_authenticated_http_subprocess_execution(self):
        import urllib.request
        import urllib.error
        with socket.socket() as sock:
            sock.bind(("127.0.0.1", 0)); port = sock.getsockname()[1]
        environment = dict(os.environ, ADAPTER_TOKEN="test-only-random-token", PORT=str(port))
        process = subprocess.Popen([sys.executable, str(Path(adapter.__file__))], env=environment, stdout=subprocess.DEVNULL, stderr=subprocess.PIPE)
        base = f"http://127.0.0.1:{port}"
        try:
            for _ in range(100):
                try:
                    urllib.request.urlopen(base + "/healthz", timeout=.2); break
                except OSError:
                    time.sleep(.05)
            body = json.dumps(request()).encode()
            with self.assertRaises(urllib.error.HTTPError) as error:
                urllib.request.urlopen(urllib.request.Request(base + "/v1/execute", data=body))
            self.assertEqual(error.exception.code, 401)
            message = urllib.request.Request(base + "/v1/execute", data=body, headers={"Authorization": "Bearer test-only-random-token", "Content-Type": "application/json"})
            result = json.load(urllib.request.urlopen(message, timeout=10))
            self.assertEqual(result["status"], "ready", result)
            self.assertEqual(result["markdown"], "Hello CloudDrive\n")
        finally:
            process.terminate(); process.communicate(timeout=5)


if __name__ == "__main__":
    unittest.main(verbosity=2)
