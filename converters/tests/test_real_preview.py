"""Optional real kkFileView 5.0.3 / LibreOffice / raster HTTP integration."""
import base64
import io
import json
import os
from pathlib import Path
import socket
import subprocess
import sys
import tempfile
import time
import unittest
import urllib.request

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import adapter


def port():
    with socket.socket() as sock:
        sock.bind(("127.0.0.1", 0))
        return sock.getsockname()[1]


def wait_http(url, process):
    for _ in range(600):
        if process.poll() is not None:
            raise AssertionError("Provider exited before readiness; inspect native provider test log")
        try:
            urllib.request.urlopen(url, timeout=.2)
            return
        except OSError:
            time.sleep(.1)
    raise AssertionError("Provider did not become ready")


@unittest.skipUnless(os.environ.get("KKFILEVIEW_JAR") and os.environ.get("KKFILEVIEW_APP_CONFIG"), "real kkFileView jar and upstream configuration not configured")
class RealPreviewTests(unittest.TestCase):
    def test_real_office_to_pdf_to_raster_http(self):
        kk_port, adapter_port, office_port = port(), port(), port()
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory); (root / "previews").mkdir(); (root / "bin").mkdir(); (root / "config").mkdir()
            config = root / "config/application.properties"
            security = Path(__file__).resolve().parents[2] / "deploy/compose/kkfileview/kkfileview.properties"
            office_home = os.environ.get("KKFILEVIEW_OFFICE_HOME", "/usr/lib/libreoffice")
            config.write_text(Path(os.environ["KKFILEVIEW_APP_CONFIG"]).read_text() + "\n" + security.read_text() + f"\nserver.port={kk_port}\nfile.dir={root / 'previews'}/\ntrust.host=127.0.0.1\noffice.plugin.server.ports={office_port}\noffice.home={office_home}\n")
            kk_log = open(os.environ.get("KKFILEVIEW_TEST_LOG", "/tmp/clouddrive-kkfileview-test.log"), "w")
            environment = dict(os.environ, HOME=str(root), KK_OFFICE_HOME=office_home, KKFILEVIEW_BIN_FOLDER=str(root / "bin"))
            kk = subprocess.Popen(["java", "-Xmx1536m", "-Djava.awt.headless=true", "-Dspring.config.location=" + str(config), "-jar", os.environ["KKFILEVIEW_JAR"]], env=environment, stdout=kk_log, stderr=subprocess.STDOUT, start_new_session=True)
            preview = None
            try:
                wait_http(f"http://127.0.0.1:{kk_port}/actuator/health", kk)
                environment.update(ADAPTER_MODE="preview", PORT=str(adapter_port), ADAPTER_TOKEN="synthetic-preview-test-only",
                                   KKFILEVIEW_URL=f"http://127.0.0.1:{kk_port}", PREVIEW_INTERNAL_URL=f"http://127.0.0.1:{adapter_port}", PREVIEW_SOURCE_DIR=str(root / "sources"))
                preview = subprocess.Popen([sys.executable, adapter.__file__], env=environment, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, start_new_session=True)
                wait_http(f"http://127.0.0.1:{adapter_port}/healthz", preview)
                from docx import Document
                from openpyxl import Workbook
                from pptx import Presentation
                doc = Document(); doc.add_heading("Private raster preview acceptance", 0); doc.add_paragraph("No source URL is exposed.")
                sheet = Workbook(); sheet.active.append(["Quarter", "Value"]); sheet.active.append(["Q1", 42])
                slides = Presentation(); slide = slides.slides.add_slide(slides.slide_layouts[1]); slide.shapes.title.text = "Controlled preview"; slide.placeholders[1].text = "Raster output only"
                for ext, document in (("docx", doc), ("xlsx", sheet), ("pptx", slides)):
                    with self.subTest(format=ext):
                        data = io.BytesIO(); document.save(data)
                        request = {"job_id": "native-preview", "attempt_id": "1", "output_scope": "tmp/native-preview/1/", "input": {"filename": "sample." + ext, "content_base64": base64.b64encode(data.getvalue()).decode()}, "recipe": {"max_pages": 10, "timeout_seconds": 120}}
                        message = urllib.request.Request(f"http://127.0.0.1:{adapter_port}/v1/execute", data=json.dumps(request).encode(), headers={"Authorization": "Bearer synthetic-preview-test-only", "Content-Type": "application/json"})
                        result = json.load(urllib.request.urlopen(message, timeout=130))
                        self.assertEqual(result["status"], "ready", result)
                        self.assertEqual(result["engine"]["name"], "kkfileview-raster-preview")
                        self.assertEqual(result["markdown"], "")
                        self.assertGreater(len(result["assets"]), 0)
                        for item in result["assets"]:
                            self.assertEqual(item["media_type"], "image/png")
                            self.assertTrue(base64.b64decode(item["content_base64"]).startswith(b"\x89PNG\r\n\x1a\n"))
                        self.assertEqual(list((root / "sources").iterdir()), [])
            finally:
                import signal
                for process in (preview, kk):
                    if process is not None:
                        os.killpg(process.pid, signal.SIGTERM)
                        try:
                            process.wait(timeout=10)
                        except subprocess.TimeoutExpired:
                            os.killpg(process.pid, signal.SIGKILL); process.wait(timeout=5)
                kk_log.close()


if __name__ == "__main__":
    unittest.main(verbosity=2)
