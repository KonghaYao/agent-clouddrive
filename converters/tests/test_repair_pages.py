"""Correctness tests for B11/B18/B19/B20/B21 and the v2 page protocol (B12).

The original audit script recorded the defects; these tests assert the fixed
behavior. OCR, vision and kkFileView endpoints are explicit local stubs. Real
HTTP, subprocess timeouts, SIGKILL, RLIMIT and filesystem paths do execute.
No external model or production service is contacted.
"""
import base64
import io
import json
import os
from pathlib import Path
import select
import shutil
import signal
import socket
import subprocess
import sys
import tempfile
import threading
import time
import unittest
import urllib.error
import urllib.parse
import urllib.request
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from unittest import mock

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import adapter
from PIL import Image, ImageDraw
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas

TOKEN = "repair-test-token"


def request(data, filename, **recipe):
    return {"job_id": "repair-job", "attempt_id": "1", "output_scope": "tmp/repair-job/1/",
            "input": {"filename": filename, "content_base64": base64.b64encode(data).decode(),
                      "sha256": __import__("hashlib").sha256(data).hexdigest()},
            "recipe": recipe}


class Stub:
    """Local HTTP stub; every received request is recorded for assertions."""

    def __init__(self, handler):
        self.requests = []
        stub = self

        class Handler(BaseHTTPRequestHandler):
            def log_message(self, *args):
                pass

            def handle_request(self):
                length = int(self.headers.get("Content-Length") or 0)
                body = self.rfile.read(length) if length else b""
                stub.requests.append({"method": self.command, "path": self.path,
                                      "headers": dict(self.headers), "body": body})
                try:
                    status, payload, content_type = handler(self.command, self.path, body)
                    self.send_response(status)
                    self.send_header("Content-Type", content_type)
                    self.send_header("Content-Length", str(len(payload)))
                    self.end_headers()
                    self.wfile.write(payload)
                except OSError:
                    pass

            do_GET = handle_request
            do_POST = handle_request

        self.server = ThreadingHTTPServer(("127.0.0.1", 0), Handler)
        self.server.daemon_threads = True
        self.thread = threading.Thread(target=self.server.serve_forever, daemon=True)
        self.thread.start()

    @property
    def url(self):
        return f"http://127.0.0.1:{self.server.server_port}"

    def close(self):
        self.server.shutdown()
        self.server.server_close()


def ocr_handler(calls, delay_second=0.0, fail_second=False):
    def handle(method, path, body):
        payload = json.loads(body or b"{}")
        calls.append(payload)
        index = len(calls)
        if delay_second and index == 2:
            time.sleep(delay_second)
        if fail_second and index == 2:
            return 500, b"{}", "application/json"
        return 200, json.dumps({"status": "ready", "markdown": f"OCR text page {index}",
                                "errors": []}).encode(), "application/json"
    return handle


def vision_handler(calls, status=200, content="Chart description"):
    def handle(method, path, body):
        payload = json.loads(body or b"{}")
        calls.append(payload)
        if status != 200:
            return status, json.dumps({"error": {"message": f"model {payload.get('model')} not found",
                                                 "type": "invalid_request_error"}}).encode(), "application/json"
        return 200, json.dumps({"choices": [{"message": {"content": content}}]}).encode(), "application/json"
    return handle


def kk_handler(pdf_bytes, source_fetches, delay=0.0):
    def handle(method, path, body):
        parsed = urllib.parse.urlsplit(path)
        if parsed.path == "/onlinePreview":
            source_url = base64.b64decode(urllib.parse.parse_qs(parsed.query)["url"][0]).decode()
            with urllib.request.urlopen(source_url, timeout=5) as response:
                source_fetches.append(response.read())
            if delay:
                time.sleep(delay)
            return 200, b"<script>var url = '/preview/result.pdf';</script>", "text/html"
        if parsed.path == "/preview/result.pdf":
            return 200, pdf_bytes, "application/pdf"
        return 404, b"not found", "text/plain"
    return handle


def empty_pdf():
    out = io.BytesIO()
    document = canvas.Canvas(out, pagesize=(80, 80))
    document.showPage()
    document.save()
    return out.getvalue()


def text_pdf(pages=1):
    out = io.BytesIO()
    document = canvas.Canvas(out)
    for page in range(1, pages + 1):
        document.drawString(50, 750, f"CloudDrive plain text report page {page}. No drawings.")
        document.showPage()
    document.save()
    return out.getvalue()


def drawing_pdf(kind):
    out = io.BytesIO()
    document = canvas.Canvas(out)
    document.drawString(40, 760, "Revenue trend chart 2024-2026 with axis labels")
    if kind == "lines":
        document.lines([(40, 100, 140, 300), (140, 300, 240, 200), (240, 200, 340, 400)])
    elif kind == "curve":
        path = document.beginPath()
        path.moveTo(40, 100)
        path.curveTo(100, 400, 200, 50, 340, 300)
        document.drawPath(path, stroke=1, fill=0)
    elif kind == "polygon":
        path = document.beginPath()
        path.moveTo(40, 100)
        path.lineTo(140, 400)
        path.lineTo(340, 150)
        path.close()
        document.drawPath(path, stroke=1, fill=1)
    else:
        raise AssertionError(kind)
    document.showPage()
    document.save()
    return out.getvalue()


def scanned_pdf(pages=1):
    out = io.BytesIO()
    document = canvas.Canvas(out)
    for _ in range(pages):
        image = Image.new("RGB", (300, 120), "white")
        ImageDraw.Draw(image).text((10, 10), "scanned page", fill="black")
        document.drawImage(ImageReader(image), 40, 40, width=300, height=120)
        document.showPage()
    document.save()
    return out.getvalue()


def multi_frame(colors):
    frames = [Image.new("RGB", (64, 64), color) for color in colors]
    out = io.BytesIO()
    frames[0].save(out, format="TIFF", save_all=True, append_images=frames[1:])
    return out.getvalue()


def docx_with_noise(pixels=300, sentence="Large legal Office preview input."):
    from docx import Document
    image = Image.frombytes("RGB", (pixels, pixels), os.urandom(pixels * pixels * 3))
    buffer = io.BytesIO()
    image.save(buffer, "PNG")
    document = Document()
    document.add_paragraph(sentence + " x" * 500)
    buffer.seek(0)
    document.add_picture(buffer)
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def vision_env(stub, model="vision-test", **extra):
    values = {"DRIVE_VISION_BASE_URL": stub.url + "/v1", "DRIVE_VISION_MODEL": model,
              "DRIVE_VISION_API_KEY": "", "DRIVE_VISION_ALLOW_EXTERNAL": "false"}
    values.update(extra)
    return mock.patch.dict(os.environ, values)


def http_request(base, data, filename, timeout=20, page_report=False, **recipe):
    body = json.dumps(request(data, filename, timeout_seconds=timeout, page_report=page_report, **recipe)).encode()
    message = urllib.request.Request(base + "/v1/execute", data=body,
                                     headers={"Authorization": "Bearer " + TOKEN, "Content-Type": "application/json"})
    with urllib.request.urlopen(message, timeout=timeout + 30) as response:
        payload = response.read()
        if response.headers.get("Content-Type", "").startswith("application/x-ndjson"):
            return [json.loads(line) for line in payload.splitlines() if line.strip()]
        return json.loads(payload)


def start_adapter(environment=None, mode="preview"):
    for _ in range(3):
        with socket.socket() as sock:
            sock.bind(("127.0.0.1", 0))
            port = sock.getsockname()[1]
        base = f"http://127.0.0.1:{port}"
        values = dict(os.environ, ADAPTER_TOKEN=TOKEN, PORT=str(port), ADAPTER_MODE=mode, PREVIEW_INTERNAL_URL=base)
        values.update(environment or {})
        process = subprocess.Popen([sys.executable, str(Path(adapter.__file__))], env=values,
                                   stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, start_new_session=True)
        for _ in range(300):
            if process.poll() is not None:
                break
            try:
                urllib.request.urlopen(base + "/healthz", timeout=.2)
                return process, base
            except OSError:
                time.sleep(.05)
        if process.poll() is None:
            stop_adapter(process)
    raise AssertionError("adapter did not become ready")


def stop_adapter(process):
    if process.poll() is None:
        os.killpg(process.pid, signal.SIGKILL)
    process.wait(timeout=5)


class MultipageImageTests(unittest.TestCase):
    """B11: multi-frame TIFF/GIF frames are real pages, never silently dropped."""

    def test_every_tiff_frame_is_converted(self):
        colors = ["red", "lime", "blue"]
        data = multi_frame(colors)
        with mock.patch.object(adapter, "ocr_remote", side_effect=["First page text", "Second page text", "Third page text"]) as ocr, \
             mock.patch.object(adapter, "describe", side_effect=["First page description", "Second page description", "Third page description"]) as vision:
            result = adapter.execute(request(data, "scanned-contract.tiff"))
        self.assertEqual(result["status"], "ready", result)
        self.assertEqual(result["completeness"]["total_pages"], 3)
        self.assertEqual([page["page"] for page in result["pages"]], [1, 2, 3])
        self.assertEqual([item["page"] for item in result["assets"]], [1, 2, 3])
        self.assertEqual((ocr.call_count, vision.call_count), (3, 3))
        self.assertEqual(result["completeness"]["missing_text_pages"], [])
        self.assertEqual(result["completeness"]["missing_description_pages"], [])
        markdown = result["markdown"]
        for page in (1, 2, 3):
            self.assertIn(f"<!-- page:{page} ", markdown)
            self.assertIn(f"![Page {page}]", markdown)
        self.assertLess(markdown.index("First page text"), markdown.index("Second page text"))
        self.assertLess(markdown.index("Second page description"), markdown.index("Third page text"))
        for item, color in zip(result["assets"], colors):
            with Image.open(io.BytesIO(base64.b64decode(item["content_base64"]))) as rendered:
                self.assertEqual(rendered.format, "PNG")
                expected = Image.new("RGB", (1, 1), color).getpixel((0, 0))
                self.assertEqual(rendered.getpixel((5, 5)), expected)

    def test_failed_second_frame_is_partial_and_not_ready(self):
        data = multi_frame(["red", "lime", "blue"])
        with mock.patch.object(adapter, "ocr_remote", side_effect=["one", adapter.AdapterError("ocr_failed", "ocr", retryable=True), "three"]), \
             mock.patch.object(adapter, "describe", return_value="visible page"):
            result = adapter.execute(request(data, "scanned-contract.tiff"))
        self.assertEqual(result["status"], "partial", result)
        self.assertEqual(result["completeness"]["status"], "partial")
        self.assertEqual(result["completeness"]["total_pages"], 3)
        self.assertEqual(result["completeness"]["missing_text_pages"], [2])
        self.assertEqual(result["completeness"]["missing_description_pages"], [])
        self.assertEqual([page["text_status"] for page in result["pages"]], ["ready", "failed", "ready"])
        self.assertEqual(len(result["assets"]), 3)
        self.assertIn("three", result["markdown"])
        self.assertNotIn("two", result["markdown"])

    def test_preview_contains_every_frame(self):
        data = multi_frame(["red", "lime", "blue"])
        result = adapter.execute(request(data, "scanned-contract.tiff"), mode="preview")
        self.assertEqual(result["status"], "ready", result)
        self.assertEqual([item["page"] for item in result["assets"]], [1, 2, 3])
        self.assertEqual(len(result["pages"]), 3)
        self.assertEqual(result["markdown"], "")

    def test_page_limit_rejects_instead_of_truncating(self):
        data = multi_frame(["red", "lime", "blue"])
        result = adapter.execute(request(data, "scanned-contract.tiff", max_pages=2))
        self.assertEqual(result["status"], "not_converted", result)
        self.assertEqual(result["errors"][0]["code"], "limit_exceeded")
        self.assertEqual(result["assets"], [])

    def test_page_selection_keeps_global_number(self):
        data = multi_frame(["red", "lime", "blue"])
        with mock.patch.object(adapter, "ocr_remote", return_value="only selected page text") as ocr, \
             mock.patch.object(adapter, "describe", return_value="selected description"):
            result = adapter.execute(request(data, "scanned-contract.tiff", page=2))
        self.assertEqual(ocr.call_count, 1)
        self.assertEqual([page["page"] for page in result["pages"]], [2])
        self.assertEqual(result["completeness"]["total_pages"], 3)
        self.assertEqual(result["completeness"]["not_processed_pages"], [1, 3])
        self.assertNotIn("page:1", result["markdown"])
        self.assertNotIn("page:3", result["markdown"])


class PdfVisualDetectionTests(unittest.TestCase):
    """B18: straight-line/vector drawings must reach the vision stage."""

    def test_content_scanner_is_text_safe_and_conservative(self):
        self.assertFalse(adapter.content_has_visual(b""))
        self.assertFalse(adapter.content_has_visual(b"BT /F1 12 Tf (10 10 m 20 20 l S) Tj ET"))
        self.assertFalse(adapter.content_has_visual(b"% 10 10 m 20 20 l S\nBT (plain title) Tj ET"))
        self.assertFalse(adapter.content_has_visual(b"BT <6d 6c 20 53> Tj ET"))
        self.assertTrue(adapter.content_has_visual(b"10 10 m 20 20 l S"))
        self.assertTrue(adapter.content_has_visual(b"0 0 10 10 re f"))
        self.assertTrue(adapter.content_has_visual(b"q /Fm1 Do Q"))
        self.assertTrue(adapter.content_has_visual(b"BI /W 1 /H 1 /BPC 8 /CS /G ID \x00 EI"))

    def test_straight_line_chart_is_sent_to_vision(self):
        calls = []
        stub = Stub(vision_handler(calls))
        try:
            with vision_env(stub):
                result = adapter.execute(request(drawing_pdf("lines"), "trend.pdf"))
        finally:
            stub.close()
        self.assertEqual(result["status"], "ready", result)
        self.assertEqual(len(calls), 1)
        self.assertEqual(calls[0]["model"], "vision-test")
        self.assertTrue(calls[0]["messages"][0]["content"][1]["image_url"]["url"].startswith("data:image/png;base64,"))
        self.assertEqual(len(result["assets"]), 1)
        self.assertEqual(result["pages"][0]["description_status"], "ready")
        self.assertIn("Chart description", result["markdown"])
        self.assertEqual(result["model"], {"provider": "openai-compatible", "name": "vision-test", "source": "deployment"})

    def test_straight_line_chart_without_vision_is_partial(self):
        with mock.patch.dict(os.environ, {"DRIVE_VISION_BASE_URL": "", "DRIVE_VISION_MODEL": ""}):
            result = adapter.execute(request(drawing_pdf("lines"), "trend.pdf"))
        self.assertEqual(result["status"], "partial", result)
        self.assertEqual(result["completeness"]["status"], "partial")
        self.assertEqual(result["completeness"]["missing_description_pages"], [1])
        self.assertEqual(result["pages"][0]["description_status"], "failed")
        self.assertEqual(len(result["assets"]), 1)
        self.assertIn("provider_not_configured", [error["code"] for error in result["errors"]])
        self.assertIn("Revenue trend chart", result["markdown"])

    def test_curve_and_filled_polygon_are_detected(self):
        for kind in ("curve", "polygon"):
            with self.subTest(kind=kind), mock.patch.object(adapter, "describe", return_value="vector drawing description") as vision:
                result = adapter.execute(request(drawing_pdf(kind), "drawing.pdf"))
            self.assertEqual(vision.call_count, 1)
            self.assertEqual(len(result["assets"]), 1, result)
            self.assertEqual(result["pages"][0]["description_status"], "ready")
            self.assertIn("vector drawing description", result["markdown"])

    def test_text_only_pdf_skips_vision(self):
        with mock.patch.object(adapter, "describe", return_value="must not run") as vision:
            result = adapter.execute(request(text_pdf(2), "plain.pdf"))
        self.assertFalse(vision.called)
        self.assertEqual(result["status"], "ready", result)
        self.assertEqual(result["assets"], [])
        self.assertEqual([page["description_status"] for page in result["pages"]], ["not_required", "not_required"])

    def test_vision_failure_is_partial_not_ready(self):
        calls = []
        stub = Stub(vision_handler(calls, status=500))
        try:
            with vision_env(stub):
                result = adapter.execute(request(drawing_pdf("lines"), "trend.pdf"))
        finally:
            stub.close()
        self.assertEqual(result["status"], "partial", result)
        self.assertEqual(result["completeness"]["missing_description_pages"], [1])
        self.assertEqual(result["errors"][0]["retryable"], True)


class RecipeModelTests(unittest.TestCase):
    """B19: the immutable recipe model executes; no silent deployment substitution."""

    def test_recipe_model_is_sent_verbatim(self):
        calls = []
        stub = Stub(vision_handler(calls))
        try:
            with vision_env(stub, model="new-deployment-model"):
                result = adapter.execute(request(drawing_pdf("lines"), "trend.pdf", model="previous-approved-model"))
        finally:
            stub.close()
        self.assertEqual(len(calls), 1)
        self.assertEqual(calls[0]["model"], "previous-approved-model")
        self.assertEqual(result["status"], "ready", result)
        self.assertEqual(result["model"], {"provider": "openai-compatible", "name": "previous-approved-model", "source": "recipe"})

    def test_provider_rejecting_recipe_model_blocks_without_substitution(self):
        calls = []
        stub = Stub(vision_handler(calls, status=404))
        try:
            with vision_env(stub, model="new-deployment-model"):
                result = adapter.execute(request(drawing_pdf("lines"), "trend.pdf", model="previous-approved-model"))
        finally:
            stub.close()
        self.assertEqual(result["status"], "blocked", result)
        self.assertEqual(result["completeness"]["status"], "blocked")
        self.assertEqual([error["code"] for error in result["errors"]], ["model_unavailable"])
        self.assertEqual(result["completeness"]["missing_description_pages"], [1])
        self.assertEqual(len(calls), 1)
        self.assertEqual(calls[0]["model"], "previous-approved-model")
        # No model executed, so provenance must not claim the recipe model ran.
        self.assertEqual(result["model"]["name"], "")

    def test_deployment_model_used_only_without_recipe_model(self):
        calls = []
        stub = Stub(vision_handler(calls))
        try:
            with vision_env(stub, model="new-deployment-model"):
                result = adapter.execute(request(drawing_pdf("lines"), "trend.pdf"))
        finally:
            stub.close()
        self.assertEqual(calls[0]["model"], "new-deployment-model")
        self.assertEqual(result["model"], {"provider": "openai-compatible", "name": "new-deployment-model", "source": "deployment"})

    def test_describe_maps_provider_404_to_model_unavailable(self):
        calls = []
        stub = Stub(vision_handler(calls, status=404))
        try:
            with vision_env(stub, model="new-deployment-model"):
                with self.assertRaises(adapter.AdapterError) as caught:
                    adapter.describe(b"png", request(drawing_pdf("lines"), "trend.pdf",
                                                    model="previous-approved-model", timeout_seconds=5))
        finally:
            stub.close()
        self.assertEqual(caught.exception.code, "model_unavailable")
        self.assertFalse(caught.exception.retryable)

    def test_recipe_model_without_provider_is_not_substituted(self):
        with mock.patch.dict(os.environ, {"DRIVE_VISION_BASE_URL": "", "DRIVE_VISION_MODEL": "new-deployment-model"}):
            with self.assertRaises(adapter.AdapterError) as caught:
                adapter.describe(b"png", request(drawing_pdf("lines"), "trend.pdf", model="previous-approved-model"))
        self.assertEqual(caught.exception.code, "provider_not_configured")


class TempRecoveryTests(unittest.TestCase):
    """B20: no staged Office original survives a timeout or SIGKILL."""

    def test_two_real_office_timeouts_leave_no_source_files(self):
        fetched = []
        stub = Stub(kk_handler(empty_pdf(), fetched, delay=10))
        with tempfile.TemporaryDirectory() as directory:
            sources = Path(directory) / "sources"
            sources.mkdir()
            process, base = start_adapter({"KKFILEVIEW_URL": stub.url, "PREVIEW_SOURCE_DIR": str(sources)})
            try:
                data = docx_with_noise(120)
                for attempt in range(2):
                    result = http_request(base, data, "document.docx", timeout=1)
                    self.assertEqual(result["status"], "failed", result)
                    self.assertEqual(result["errors"][0]["code"], "timeout", result)
                    self.assertEqual(list(sources.iterdir()), [], f"attempt {attempt} left temporary originals")
                self.assertEqual(len(fetched), 2)
                self.assertEqual(fetched[0], data)
            finally:
                stop_adapter(process)
        stub.close()

    def test_sigkill_orphan_is_reaped_on_next_start(self):
        fetched = []
        stub = Stub(kk_handler(empty_pdf(), fetched, delay=10))
        with tempfile.TemporaryDirectory() as directory:
            sources = Path(directory) / "sources"
            sources.mkdir()
            process, base = start_adapter({"KKFILEVIEW_URL": stub.url, "PREVIEW_SOURCE_DIR": str(sources)})
            outcome = {}

            def call():
                try:
                    outcome["result"] = http_request(base, docx_with_noise(120), "document.docx", timeout=2)
                except Exception as error:  # the server is killed mid-request.
                    outcome["error"] = error

            thread = threading.Thread(target=call, daemon=True)
            thread.start()
            orphan = None
            deadline = time.time() + 10
            while time.time() < deadline:
                for candidate in sources.iterdir():
                    if candidate.name.startswith("attempt-") and any(item.suffix == ".docx" for item in candidate.iterdir()):
                        orphan = candidate
                        break
                if orphan is not None:
                    break
                time.sleep(.05)
            self.assertIsNotNone(orphan, "the original was never staged")
            stop_adapter(process)
            self.assertTrue(orphan.is_dir(), "attempt directory vanished before reaping")
            thread.join(timeout=5)
            # The detached child exits on its own timeout; only the parent dir remains.
            time.sleep(3.5)
            process2, base2 = start_adapter({"KKFILEVIEW_URL": stub.url, "PREVIEW_SOURCE_DIR": str(sources)})
            try:
                deadline = time.time() + 5
                while time.time() < deadline and orphan.exists():
                    time.sleep(.05)
                self.assertFalse(orphan.exists(), "startup reaping did not remove the SIGKILL orphan")
            finally:
                stop_adapter(process2)
        stub.close()

    def test_reaper_keeps_active_and_foreign_directories(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            active = adapter.AttemptDirectory.create(root=root)
            marker = json.loads((active.path / adapter.ATTEMPT_MARKER).read_text())
            marker["created"] = 0  # stale by TTL, but the flock proves it is in use
            (active.path / adapter.ATTEMPT_MARKER).write_text(json.dumps(marker))
            foreign = root / ("attempt-" + "a" * 32)
            foreign.mkdir()
            (foreign / adapter.ATTEMPT_MARKER).write_text(json.dumps({"service": "other-service", "pid": 999999, "created": 0}))
            plain = root / "preview-data"
            plain.mkdir()
            dead = root / ("attempt-" + "b" * 32)
            dead.mkdir()
            (dead / adapter.ATTEMPT_MARKER).write_text(json.dumps({"service": adapter.ATTEMPT_SERVICE, "pid": 999999, "created": time.time()}))
            stale = adapter.AttemptDirectory.create(root=root)
            stale.close()
            marker = json.loads((stale.path / adapter.ATTEMPT_MARKER).read_text())
            marker["created"] = time.time() - 10000
            (stale.path / adapter.ATTEMPT_MARKER).write_text(json.dumps(marker))
            try:
                summary = adapter.reap_orphans(root=root, ttl=1800)
                self.assertTrue(active.path.is_dir(), "active attempt was deleted")
                self.assertTrue(foreign.is_dir(), "foreign service directory was deleted")
                self.assertTrue(plain.is_dir(), "non-attempt directory was deleted")
                self.assertFalse(dead.exists(), "owner-dead orphan was not reaped")
                self.assertFalse(stale.path.exists(), "TTL-expired orphan was not reaped")
                self.assertIn(dead.name, summary["removed"])
                self.assertIn(stale.path.name, summary["removed"])
                self.assertIn(active.name, summary["active"])
                self.assertIn(foreign.name, summary["foreign"])
            finally:
                active.close()

    def test_cleanup_failure_is_diagnosed(self):
        root = tempfile.mkdtemp(prefix="repair-cleanup-")
        try:
            attempt = adapter.AttemptDirectory.create(root=root)
            result = adapter.empty_result()
            with mock.patch("adapter.shutil.rmtree", side_effect=OSError("resource busy")):
                diagnostic = adapter.cleanup_attempt(attempt, root=root)
            self.assertEqual(diagnostic["code"], "temp_cleanup_failed")
            adapter.attach_cleanup_warning(result, diagnostic)
            self.assertIn("temp_cleanup_failed", result["completeness"]["warnings"])
            log = (Path(root) / "cleanup-diagnostics.log").read_text()
            self.assertIn("temp_cleanup_failed", log)
            self.assertTrue((Path(root) / attempt.name).is_dir())
        finally:
            shutil.rmtree(root, ignore_errors=True)


class TempBudgetTests(unittest.TestCase):
    """B21: input staging, output and total temporary space are separate budgets."""

    def test_legal_large_input_with_tiny_output_cap_is_previewed(self):
        data = docx_with_noise(300)
        self.assertGreater(len(data), 200_000)
        fetched = []
        stub = Stub(kk_handler(empty_pdf(), fetched))
        with tempfile.TemporaryDirectory() as directory:
            sources = Path(directory) / "sources"
            sources.mkdir()
            process, base = start_adapter({"KKFILEVIEW_URL": stub.url, "PREVIEW_SOURCE_DIR": str(sources),
                                           "ADAPTER_MAX_INPUT_BYTES": str(2 * 1024 * 1024),
                                           "ADAPTER_MAX_OUTPUT_BYTES": "4096",
                                           "ADAPTER_MAX_TEMP_BYTES": str(8 * 1024 * 1024)})
            try:
                result = http_request(base, data, "document.docx", timeout=20)
            finally:
                stop_adapter(process)
                stub.close()
            self.assertEqual(result["status"], "ready", result)
            self.assertEqual(result["engine"]["name"], "kkfileview-raster-preview")
            self.assertEqual(len(fetched), 1)
            self.assertEqual(fetched[0], data)  # the whole original was staged, not truncated
            self.assertGreaterEqual(len(result["assets"]), 1)
            self.assertEqual(list(sources.iterdir()), [])

    def test_output_budget_exceeded_fails_and_cleans(self):
        data = docx_with_noise(300)
        fetched = []
        stub = Stub(kk_handler(empty_pdf(), fetched))
        with tempfile.TemporaryDirectory() as directory:
            sources = Path(directory) / "sources"
            sources.mkdir()
            process, base = start_adapter({"KKFILEVIEW_URL": stub.url, "PREVIEW_SOURCE_DIR": str(sources),
                                           "ADAPTER_MAX_INPUT_BYTES": str(2 * 1024 * 1024),
                                           "ADAPTER_MAX_OUTPUT_BYTES": "4096",
                                           "ADAPTER_MAX_TEMP_BYTES": str(8 * 1024 * 1024)})
            try:
                result = http_request(base, data, "document.docx", timeout=20, max_output_bytes=64)
            finally:
                stop_adapter(process)
                stub.close()
            self.assertEqual(result["status"], "failed", result)
            self.assertEqual(result["errors"][0]["code"], "resource_limit")
            self.assertEqual(result["errors"][0]["stage"], "output")
            self.assertEqual(list(sources.iterdir()), [])

    def test_input_over_limit_is_rejected_without_staging(self):
        data = docx_with_noise(160)
        self.assertGreater(len(data), 64 * 1024)
        fetched = []
        stub = Stub(kk_handler(empty_pdf(), fetched))
        with tempfile.TemporaryDirectory() as directory:
            sources = Path(directory) / "sources"
            sources.mkdir()
            process, base = start_adapter({"KKFILEVIEW_URL": stub.url, "PREVIEW_SOURCE_DIR": str(sources),
                                           "ADAPTER_MAX_INPUT_BYTES": str(64 * 1024)})
            try:
                result = http_request(base, data, "document.docx", timeout=20)
            finally:
                stop_adapter(process)
                stub.close()
            self.assertEqual(result["status"], "not_converted", result)
            self.assertEqual(result["errors"][0]["code"], "limit_exceeded")
            self.assertEqual(list(sources.iterdir()), [])
            self.assertEqual(fetched, [])

    def test_total_temp_budget_stops_staging_and_cleans(self):
        data = docx_with_noise(300)
        fetched = []
        stub = Stub(kk_handler(empty_pdf(), fetched))
        with tempfile.TemporaryDirectory() as directory:
            sources = Path(directory) / "sources"
            sources.mkdir()
            process, base = start_adapter({"KKFILEVIEW_URL": stub.url, "PREVIEW_SOURCE_DIR": str(sources),
                                           "ADAPTER_MAX_INPUT_BYTES": str(2 * 1024 * 1024),
                                           "ADAPTER_MAX_OUTPUT_BYTES": "4096",
                                           "ADAPTER_MAX_TEMP_BYTES": str(512 * 1024)})
            try:
                result = http_request(base, data, "document.docx", timeout=20)
            finally:
                stop_adapter(process)
                stub.close()
            self.assertEqual(result["status"], "failed", result)
            self.assertEqual(result["errors"][0]["code"], "temp_budget_exceeded")
            self.assertEqual(list(sources.iterdir()), [])
            self.assertEqual(fetched, [])

    def test_staging_budget_and_write_failure_leave_no_fragments(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            with self.assertRaises(adapter.AdapterError) as caught:
                adapter.stage_source(root, b"x" * 4096, "docx", 1024)
            self.assertEqual(caught.exception.code, "temp_budget_exceeded")
            self.assertEqual(list(root.iterdir()), [])
            with mock.patch("adapter.os.fsync", side_effect=OSError(28, "No space left on device")):
                with self.assertRaises(adapter.AdapterError) as caught:
                    adapter.stage_source(root, b"x" * 1024, "docx", 4096)
            self.assertEqual(caught.exception.code, "temp_write_failed")
            self.assertEqual(list(root.iterdir()), [])
            with mock.patch("adapter.os.fsync", side_effect=OSError(28, "No space left on device")):
                with self.assertRaises(adapter.AdapterError) as caught:
                    adapter.write_bounded_file(root / "input.json", b"x" * 10, 100)
            self.assertEqual(caught.exception.code, "temp_write_failed")
            self.assertEqual(list(root.iterdir()), [])

    def test_staged_source_is_private(self):
        with tempfile.TemporaryDirectory() as directory:
            attempt = adapter.AttemptDirectory.create(root=directory)
            try:
                self.assertEqual(attempt.path.stat().st_mode & 0o777, 0o700)
                source = adapter.stage_source(attempt.path, b"PK\x03\x04test", "docx", 1024)
                self.assertEqual(source.stat().st_mode & 0o777, 0o600)
            finally:
                adapter.cleanup_attempt(attempt, root=directory)


class PageProtocolTests(unittest.TestCase):
    """B12: JSONL page lines, immediate flush, resume only failed pages."""

    def test_page_report_jsonl_order_and_incremental_flush(self):
        ocr_calls, vision_calls = [], []
        ocr = Stub(ocr_handler(ocr_calls, delay_second=2.5))
        vision = Stub(vision_handler(vision_calls))
        try:
            with tempfile.TemporaryDirectory() as directory:
                inp = Path(directory) / "input.json"
                inp.write_text(json.dumps(request(scanned_pdf(2), "scan.pdf", timeout_seconds=30, page_report=True)))
                environment = dict(os.environ, DRIVE_OCR_URL=ocr.url, DRIVE_VISION_BASE_URL=vision.url + "/v1",
                                   DRIVE_VISION_MODEL="vision-test", DRIVE_VISION_API_KEY="")
                process = subprocess.Popen([sys.executable, str(Path(adapter.__file__)), "--execute", str(inp)],
                                           stdout=subprocess.PIPE, stderr=subprocess.PIPE, env=environment)
                try:
                    ready, _, _ = select.select([process.stdout], [], [], 1.5)
                    self.assertTrue(ready, "no JSONL line was flushed before the slow second page finished")
                    first = process.stdout.readline()
                    rest, errors = process.communicate(timeout=60)
                finally:
                    if process.poll() is None:
                        process.kill()
                        process.wait()
                self.assertEqual(process.returncode, 0, errors.decode())
                lines = [first] + rest.splitlines()
                events = [json.loads(line) for line in lines if line.strip()]
        finally:
            ocr.close()
            vision.close()
        self.assertTrue(events)
        self.assertTrue(all(event.get("type") == "page" for event in events[:-1]), events)
        self.assertEqual(events[-1]["type"], "result")
        result = events[-1]
        self.assertEqual(result["status"], "ready", result)
        self.assertEqual(result["completeness"]["total_pages"], 2)
        pages = [(event["page"], event["stage"], event["state"]) for event in events[:-1]]
        self.assertEqual([page for page, _, _ in pages], [1] * 4 + [2] * 4)
        for page in (1, 2):
            stages = [stage for page_no, stage, _ in pages if page_no == page]
            self.assertEqual(stages, ["render", "ocr", "text", "description"])
        for event in events[:-1]:
            self.assertEqual(event["state"], "ready", event)
            self.assertTrue({"markdown", "assets", "pages"} <= set(event["result_ref"]))
        description_events = [event for event in events[:-1] if event["stage"] == "description"]
        self.assertTrue(description_events)
        for event in description_events:
            self.assertEqual(event["result_ref"]["model"], "vision-test")
            self.assertEqual(event["result_ref"]["provider"], "openai-compatible")
        # Replaying the flushed page slices reproduces the final page result.
        replayed_markdown = {1: "", 2: ""}
        replayed_assets = {1: [], 2: []}
        replayed_pages = {}
        for event in events[:-1]:
            reference = event["result_ref"]
            replayed_markdown[event["page"]] += reference["markdown"]
            replayed_assets[event["page"]] += reference["assets"]
            replayed_pages[event["page"]] = reference["pages"][0]
        self.assertEqual(replayed_markdown[1] + replayed_markdown[2], result["markdown"])
        self.assertEqual(replayed_assets[1] + replayed_assets[2], result["assets"])
        for page in (1, 2):
            self.assertEqual(replayed_pages[page], next(item for item in result["pages"] if item["page"] == page))
        self.assertEqual(result["engine"]["model"], "vision-test")
        self.assertEqual(result["engine"]["provider"], "openai-compatible")
        self.assertEqual(result["model"]["name"], "vision-test")
        self.assertIn("OCR text page 1", result["markdown"])
        self.assertIn("OCR text page 2", result["markdown"])

    def test_page_report_disabled_keeps_single_json(self):
        with tempfile.TemporaryDirectory() as directory:
            inp = Path(directory) / "input.json"
            inp.write_text(json.dumps(request(text_pdf(1), "plain.pdf", timeout_seconds=30)))
            environment = dict(os.environ, DRIVE_OCR_URL="", DRIVE_VISION_BASE_URL="", DRIVE_VISION_MODEL="")
            process = subprocess.run([sys.executable, str(Path(adapter.__file__)), "--execute", str(inp)],
                                     capture_output=True, timeout=60, env=environment)
        self.assertEqual(process.returncode, 0, process.stderr.decode())
        output = process.stdout.decode().strip()
        self.assertNotIn("\n", output)
        result = json.loads(output)
        self.assertNotIn("type", result)
        self.assertEqual(result["status"], "ready", result)

    def test_resume_pages_only_runs_the_failed_page(self):
        data = scanned_pdf(3)
        with mock.patch.object(adapter, "ocr_remote", side_effect=["page one text", adapter.AdapterError("ocr_failed", "ocr", retryable=True), "page three text"]), \
             mock.patch.object(adapter, "describe", return_value="page description"):
            first = adapter.execute(request(data, "scan.pdf"))
        self.assertEqual(first["status"], "partial", first)
        self.assertEqual(first["completeness"]["missing_text_pages"], [2])
        with mock.patch.object(adapter, "ocr_remote", return_value="retried page two") as retry_ocr, \
             mock.patch.object(adapter, "describe", return_value="retried description"):
            reporter = adapter.PageReporter(True, io.StringIO())
            resumed = adapter.execute(request(data, "scan.pdf", resume_pages=[2]), reporter=reporter)
        self.assertEqual(retry_ocr.call_count, 1)
        self.assertEqual([page["page"] for page in resumed["pages"]], [2])
        self.assertEqual(resumed["completeness"]["total_pages"], 3)
        self.assertEqual(resumed["completeness"]["not_processed_pages"], [1, 3])
        self.assertEqual(resumed["status"], "ready", resumed)
        self.assertIn("retried page two", resumed["markdown"])
        self.assertNotIn("page one text", resumed["markdown"])
        self.assertNotIn("page three text", resumed["markdown"])
        events = [json.loads(line) for line in reporter.stream.getvalue().splitlines()]
        self.assertTrue(events)
        self.assertEqual({event["page"] for event in events}, {2})
        self.assertEqual([event["stage"] for event in events], ["render", "ocr", "text", "description"])
        with mock.patch.object(adapter, "ocr_remote", return_value="retried page two"), \
             mock.patch.object(adapter, "describe", return_value="retried description"):
            again = adapter.execute(request(data, "scan.pdf", resume_pages=[2]))
        self.assertEqual(again["markdown"], resumed["markdown"])
        self.assertEqual(again["pages"], resumed["pages"])

    def test_resume_page_out_of_range_is_rejected(self):
        with mock.patch.object(adapter, "ocr_remote", return_value="text"), mock.patch.object(adapter, "describe", return_value="desc"):
            result = adapter.execute(request(scanned_pdf(2), "scan.pdf", resume_pages=[3]))
        self.assertEqual(result["status"], "not_converted", result)
        self.assertEqual(result["errors"][0]["code"], "limit_exceeded")

    def test_http_page_report_timeout_keeps_confirmed_pages(self):
        ocr_calls, vision_calls = [], []
        ocr = Stub(ocr_handler(ocr_calls, delay_second=10))
        vision = Stub(vision_handler(vision_calls))
        process, base = start_adapter({"DRIVE_OCR_URL": ocr.url, "DRIVE_VISION_BASE_URL": vision.url + "/v1",
                                       "DRIVE_VISION_MODEL": "vision-test"}, mode="anydoc")
        try:
            events = http_request(base, scanned_pdf(2), "scan.pdf", timeout=2, page_report=True)
        finally:
            stop_adapter(process)
            ocr.close()
            vision.close()
        self.assertTrue(events)
        self.assertEqual(events[-1]["type"], "result")
        self.assertEqual(events[-1]["status"], "failed", events[-1])
        self.assertEqual(events[-1]["errors"][0]["code"], "timeout")
        confirmed = [(event["page"], event["stage"]) for event in events[:-1] if event.get("type") == "page"]
        self.assertIn((1, "render"), confirmed)
        self.assertIn((1, "ocr"), confirmed)
        self.assertIn((1, "description"), confirmed)
        self.assertIn((2, "render"), confirmed)
        self.assertNotIn((2, "ocr"), confirmed)
        self.assertNotIn((2, "description"), confirmed)

    def test_http_page_report_streams_ndjson(self):
        process, base = start_adapter({}, mode="anydoc")
        try:
            events = http_request(base, text_pdf(2), "plain.pdf", timeout=20, page_report=True)
        finally:
            stop_adapter(process)
        self.assertTrue(events)
        self.assertTrue(all(event.get("type") == "page" for event in events[:-1]), events)
        self.assertEqual(events[-1]["type"], "result")
        self.assertEqual(events[-1]["status"], "ready", events[-1])
        self.assertEqual([event["page"] for event in events[:-1]], [1, 1, 2, 2])
        self.assertEqual([event["stage"] for event in events[:-1]], ["text", "description", "text", "description"])
        self.assertEqual([event["state"] for event in events[:-1]], ["ready", "not_required", "ready", "not_required"])
