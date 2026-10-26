#!/usr/bin/env python3
"""Generate non-private fixtures for manual converter/preview validation."""
import io
from pathlib import Path
import sys
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "tests"))
from test_adapter import pdf
from test_real_ocr import image

root = Path(__file__).resolve().parent
(root / "sample.md").write_text("# CloudDrive 验收\n\nStable source version sample.\n", encoding="utf-8")
(root / "sample.txt").write_text("CloudDrive plain text.\n", encoding="utf-8")
(root / "text.pdf").write_bytes(pdf(3))
(root / "mixed-scan.pdf").write_bytes(pdf(3, scanned=(2,)))
(root / "ocr.png").write_bytes(image())
(root / "sample.csv").write_text("name,value\nAlice,17\n", encoding="utf-8")
(root / "sample.rtf").write_bytes(br"{\rtf1\ansi CloudDrive RTF document.}")
(root / "safe.svg").write_text('<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 100 100"><rect width="100" height="100" fill="#1677ff"/><text x="10" y="55" fill="white">Drive</text></svg>')
from docx import Document
value = Document(); value.add_heading("CloudDrive sample", 0); value.add_paragraph("Immutable document source version."); value.save(root / "sample.docx")
from openpyxl import Workbook
value = Workbook(); value.active.append(["Owner", "Quota"]); value.active.append(["Alice", 200]); value.save(root / "sample.xlsx")
from pptx import Presentation
value = Presentation(); slide = value.slides.add_slide(value.slide_layouts[1]); slide.shapes.title.text = "CloudDrive sample"; slide.placeholders[1].text = "Immutable version references"; value.save(root / "sample.pptx")
print("Generated eleven synthetic public fixtures.")
